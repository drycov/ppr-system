"""Импорт объектов ИНТ/ВОЛС из Excel-отчёта ППР (кроссплатформенно)."""
from pathlib import Path
from datetime import date
from sqlalchemy.orm import Session
import openpyxl
from models import Object, PPRPlan

# Ищем файл в нескольких местах (Windows / Linux / рядом с проектом)
SEARCH_NAMES = [
    "Отчет_ППР_АПК_СЕНТЯБРЬ_2026г..xlsx",
    "Отчет_ППР_АПК_СЕНТЯБРЬ_2026г.xlsx",
    "Otchet_PPR_APK.xlsx",
    "ppr_objects.xlsx",
]


def find_report_file(explicit: str | None = None) -> Path | None:
    if explicit:
        p = Path(explicit)
        if p.exists():
            return p
    base = Path(__file__).resolve().parent
    candidates = [
        base / "data",
        base.parent / "data",
        base.parent / "attachments",
        Path("/home/workdir/attachments"),
        Path.home() / "Downloads",
        Path("D:/dev/ppr-system/data"),
        Path("D:/dev/ppr-system/backend/data"),
    ]
    for folder in candidates:
        if not folder.exists():
            continue
        for name in SEARCH_NAMES:
            f = folder / name
            if f.exists():
                return f
        # любой xlsx с ППР в имени
        for f in folder.glob("*.xlsx"):
            if "ППР" in f.name.upper() or "PPR" in f.name.upper() or "ОТЧЕТ" in f.name.upper():
                return f
    return None


def import_from_excel(db: Session, file_path: str | None = None) -> dict:
    path = find_report_file(file_path)
    if not path:
        return {
            "ok": False,
            "message": "Файл отчёта не найден. Положите Excel в backend/data/ или укажите путь.",
            "int": 0,
            "voles": 0,
        }

    wb = openpyxl.load_workbook(path, data_only=True)
    int_count = 0
    voles_count = 0

    # ИНТЕГ
    sheet_int = None
    for name in wb.sheetnames:
        if "ИНТЕГ" in name.upper() or "INTEG" in name.upper():
            sheet_int = wb[name]
            break
    if sheet_int is None and wb.sheetnames:
        sheet_int = wb[wb.sheetnames[0]]

    if sheet_int:
        for row in sheet_int.iter_rows(min_row=4, max_row=400, max_col=9, values_only=True):
            num, code, address, cameras, responsible, ppr_date, period, status, work_types = (list(row) + [None]*9)[:9]
            if not code:
                continue
            code = str(code).strip()
            if not (code.startswith("ИНТ") or code.startswith("INT")):
                continue
            obj = db.query(Object).filter(Object.code == code).first()
            try:
                cams = int(cameras) if cameras is not None and str(cameras).replace(".0","").isdigit() else 0
            except Exception:
                cams = 0
            if not obj:
                obj = Object(
                    code=code,
                    name=str(address)[:100] if address else code,
                    address=str(address) if address else None,
                    object_type="camera_integration",
                    cameras_count=cams,
                    status="active",
                    responsible_default=str(responsible).strip() if responsible else None,
                )
                db.add(obj)
                int_count += 1
            else:
                obj.address = str(address) if address else obj.address
                obj.cameras_count = cams or obj.cameras_count
                obj.responsible_default = str(responsible).strip() if responsible else obj.responsible_default

    # ВОЛС
    sheet_v = None
    for name in wb.sheetnames:
        if "ВОЛС" in name.upper() or "VOLES" in name.upper() or "ВОЛС" in name:
            sheet_v = wb[name]
            break
    if sheet_v:
        for row in sheet_v.iter_rows(min_row=4, max_row=300, max_col=7, values_only=True):
            num, address, responsible, ppr_date, period, status, note = (list(row) + [None]*7)[:7]
            if not address:
                continue
            code = f"ВОЛС-{num}" if num else f"ВОЛС-{voles_count+1}"
            if db.query(Object).filter(Object.code == code).first():
                continue
            db.add(Object(
                code=code,
                name=str(address)[:100],
                address=str(address),
                object_type="voles",
                status="active",
                responsible_default=str(responsible).strip() if responsible else None,
            ))
            voles_count += 1

    db.commit()
    return {
        "ok": True,
        "message": f"Импорт из {path.name}: ИНТ +{int_count}, ВОЛС +{voles_count}",
        "file": str(path),
        "int": int_count,
        "voles": voles_count,
        "total_objects": db.query(Object).count(),
    }
