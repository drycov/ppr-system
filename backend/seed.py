"""Инициализация БД: шаблоны из Положения ОСС + импорт объектов из Excel"""
from datetime import date, datetime, timedelta
from sqlalchemy.orm import Session
from database import SessionLocal, engine, Base
from auth import get_password_hash
from models import Object, PPRTemplate, PPRPlan, User, DailyCheck, Department, Position, SystemRole
import openpyxl
from pathlib import Path

# Путь к отчётам
ATTACHMENTS = Path("/home/workdir/attachments")


def create_tables():
    Base.metadata.create_all(bind=engine)


def seed_templates(db: Session):
    """
    Шаблоны ППР по распределению:
    - sysadmin — системные администраторы
    - installer — монтажники
    - network_engineer — сетевые инженеры
    + АПК/ВОЛС (камера / линии)
    """
    templates = [
        # ========== 1. Системные администраторы ==========
        {
            "code": "SA-1",
            "name": "Проверка серверов DNS, DHCP, мониторинга и контроллеров",
            "description": "Диски, RAID, температура серверов сетевых служб",
            "object_types": ["network"],
            "work_items": ["Проверка дисков/RAID", "Контроль температуры", "Состояние контроллеров"],
            "periodicity": "quarterly",
            "responsible_role": "engineer",
            "work_direction": "sysadmin",
        },
        {
            "code": "SA-2",
            "name": "Обновление системного ПО, очистка временных файлов",
            "description": "Обновление ПО серверов сетевых служб, очистка temp",
            "object_types": ["network"],
            "work_items": ["Обновление ПО", "Очистка temp"],
            "periodicity": "quarterly",
            "responsible_role": "engineer",
            "work_direction": "sysadmin",
        },
        {
            "code": "SA-3",
            "name": "Проверка резервных копий конфигураций",
            "description": "Проверка успешности автоматического резервирования конфигураций",
            "object_types": ["network"],
            "work_items": ["Проверка бэкапов конфигураций"],
            "periodicity": "daily",
            "responsible_role": "engineer",
            "work_direction": "sysadmin",
        },
        {
            "code": "SA-4",
            "name": "Проверка системы мониторинга (Zabbix и собств. решения)",
            "description": "Агенты, шаблоны, оповещения",
            "object_types": ["network"],
            "work_items": ["Проверка агентов", "Актуализация шаблонов", "Проверка оповещений"],
            "periodicity": "monthly",
            "responsible_role": "engineer",
            "work_direction": "sysadmin",
        },
        {
            "code": "SA-5",
            "name": "Анализ системных журналов и событий безопасности",
            "description": "Ежедневный разбор логов и security-событий",
            "object_types": ["network"],
            "work_items": ["Анализ логов", "События безопасности"],
            "periodicity": "daily",
            "responsible_role": "engineer",
            "work_direction": "sysadmin",
        },
        {
            "code": "SA-6",
            "name": "Контроль состояния ИБП по мониторингу",
            "description": "Заряд, напряжение, температура по данным мониторинга",
            "object_types": ["network"],
            "work_items": ["Заряд АКБ", "Напряжение", "Температура ИБП"],
            "periodicity": "daily",
            "responsible_role": "engineer",
            "work_direction": "sysadmin",
        },
        {
            "code": "SA-7",
            "name": "Проверка VPN-служб, туннелей и доступности сетевых сервисов",
            "description": "Сервисная проверка VPN и доступности ключевых сетевых сервисов",
            "object_types": ["network"],
            "work_items": ["VPN-службы", "Туннели", "Доступность сервисов"],
            "periodicity": "monthly",
            "responsible_role": "engineer",
            "work_direction": "sysadmin",
        },
        # ========== 2. Монтажники ==========
        {
            "code": "IN-1",
            "name": "Внешний осмотр СКС, патч-панелей и кроссов",
            "description": "Визуальный осмотр структурированной кабельной системы",
            "object_types": ["network"],
            "work_items": ["Осмотр СКС", "Патч-панели", "Кроссы"],
            "periodicity": "annual",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
        {
            "code": "IN-2",
            "name": "Проверка маркировки кабелей и портов",
            "description": "Актуальность обозначений кабелей и портов",
            "object_types": ["network"],
            "work_items": ["Маркировка кабелей", "Маркировка портов"],
            "periodicity": "annual",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
        {
            "code": "IN-3",
            "name": "Тестирование критичных кабельных линий",
            "description": "Тестирование критичных линий, устранение повреждений",
            "object_types": ["network"],
            "work_items": ["Тестирование линий", "Устранение повреждений"],
            "periodicity": "annual",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
        {
            "code": "IN-4",
            "name": "Проверка креплений сетевого оборудования",
            "description": "Проверка надёжности креплений активного оборудования",
            "object_types": ["network"],
            "work_items": ["Проверка креплений"],
            "periodicity": "quarterly",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
        {
            "code": "IN-5",
            "name": "Очистка оборудования от пыли, вентиляторы, температура",
            "description": "Очистка активного оборудования, проверка вентиляторов и температурного режима",
            "object_types": ["network"],
            "work_items": ["Очистка от пыли", "Проверка вентиляторов", "Контроль температуры"],
            "periodicity": "quarterly",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
        {
            "code": "IN-6",
            "name": "Очистка оборудования беспроводной инфраструктуры",
            "description": "Очистка точек доступа Wi-Fi и связанного оборудования",
            "object_types": ["network"],
            "work_items": ["Очистка Wi-Fi оборудования"],
            "periodicity": "quarterly",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
        {
            "code": "IN-7",
            "name": "Проверка состояния ИБП, АКБ и автономности",
            "description": "Физическая проверка ИБП, аккумуляторов и автономности (II и IV кварталы)",
            "object_types": ["network"],
            "work_items": ["Тест батарей", "Проверка автономии", "Состояние АКБ"],
            "periodicity": "semi_annual",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
        # ========== 3. Сетевые инженеры ==========
        {
            "code": "NE-1",
            "name": "Аудит конфигураций маршрутизаторов и коммутаторов",
            "description": "Аудит конфигураций, проверка резервных копий, анализ логов",
            "object_types": ["network"],
            "work_items": ["Аудит конфигураций", "Проверка бэкапов", "Анализ логов"],
            "periodicity": "monthly",
            "responsible_role": "engineer",
            "work_direction": "network_engineer",
        },
        {
            "code": "NE-2",
            "name": "Проверка VPN, туннелей, HotSpot и резервных каналов",
            "description": "Работоспособность VPN, туннелей, HotSpot, резервных каналов",
            "object_types": ["network"],
            "work_items": ["VPN", "Туннели", "HotSpot", "Резервные каналы"],
            "periodicity": "monthly",
            "responsible_role": "engineer",
            "work_direction": "network_engineer",
        },
        {
            "code": "NE-3",
            "name": "Аудит правил Firewall и событий безопасности",
            "description": "Аудит правил, события безопасности, сигнатуры",
            "object_types": ["network"],
            "work_items": ["Аудит правил Firewall", "События безопасности", "Сигнатуры"],
            "periodicity": "monthly",
            "responsible_role": "engineer",
            "work_direction": "network_engineer",
        },
        {
            "code": "NE-4",
            "name": "Проверка точек доступа Wi-Fi, покрытия и загрузки",
            "description": "Покрытие, загруженность радиоканалов",
            "object_types": ["network"],
            "work_items": ["Проверка покрытия", "Анализ загрузки каналов"],
            "periodicity": "quarterly",
            "responsible_role": "engineer",
            "work_direction": "network_engineer",
        },
        {
            "code": "NE-5",
            "name": "Проверка и безопасное обновление прошивок",
            "description": "Обновления прошивок сетевого оборудования",
            "object_types": ["network"],
            "work_items": ["Проверка обновлений", "Безопасное обновление"],
            "periodicity": "quarterly",
            "responsible_role": "engineer",
            "work_direction": "network_engineer",
        },
        {
            "code": "NE-6",
            "name": "Тестирование переключения на резервные каналы, SLA",
            "description": "Failover, проверка SLA провайдеров",
            "object_types": ["network"],
            "work_items": ["Тест failover", "Проверка SLA"],
            "periodicity": "quarterly",
            "responsible_role": "engineer",
            "work_direction": "network_engineer",
        },
        {
            "code": "NE-7",
            "name": "Контроль загрузки каналов связи и ключевых линков",
            "description": "Ежедневный контроль утилизации каналов",
            "object_types": ["network"],
            "work_items": ["Загрузка каналов", "Ключевые линки"],
            "periodicity": "daily",
            "responsible_role": "engineer",
            "work_direction": "network_engineer",
        },
        {
            "code": "NE-8",
            "name": "Проверка доступности критических узлов и сервисов",
            "description": "Доступность критических узлов, сервисов и каналов связи",
            "object_types": ["network"],
            "work_items": ["Доступность узлов", "Критические сервисы"],
            "periodicity": "daily",
            "responsible_role": "engineer",
            "work_direction": "network_engineer",
        },
        {
            "code": "NE-9",
            "name": "Инвентаризация сетевого оборудования, схемы и документация",
            "description": "Полная инвентаризация, актуализация схем (IV квартал) — ответственный начальник ОСС",
            "object_types": ["network"],
            "work_items": ["Инвентаризация", "Схемы сети", "Документация"],
            "periodicity": "annual",
            "responsible_role": "head",
            "work_direction": "head",
        },
        # ========== ЦОД — отдельный регламент эксплуатации дата-центра ==========
        {
            "code": "DC-1",
            "name": "Ежедневный контроль ЦОД",
            "description": "Контроль температуры и климатических условий, питания и ИБП, аварийных сигналов, доступности критичных сервисов",
            "object_types": ["cod"],
            "work_items": ["Температура и климат", "Питание и ИБП", "Аварийные сигналы", "Доступность критичных сервисов"],
            "periodicity": "daily",
            "responsible_role": "engineer",
            "work_direction": "sysadmin",
        },
        {
            "code": "DC-2",
            "name": "Ежемесячная проверка инженерных систем ЦОД",
            "description": "Проверка ИБП и состояния АКБ, климатического оборудования, датчиков и систем оповещения",
            "object_types": ["cod"],
            "work_items": ["Состояние ИБП и АКБ", "Климатическое оборудование", "Датчики среды", "Система оповещения"],
            "periodicity": "monthly",
            "responsible_role": "engineer",
            "work_direction": "sysadmin",
        },
        {
            "code": "DC-3",
            "name": "Квартальный осмотр инфраструктуры ЦОД",
            "description": "Проверка стоек, кабельной инфраструктуры, маркировки, заземления и документации",
            "object_types": ["cod"],
            "work_items": ["Состояние стоек", "Кабельная инфраструктура и маркировка", "Проверка заземления", "Актуальность схем и документации"],
            "periodicity": "quarterly",
            "responsible_role": "engineer",
            "work_direction": "network_engineer",
        },
        {
            "code": "DC-4",
            "name": "Полугодовая проверка безопасности ЦОД",
            "description": "Проверка контроля доступа, пожарной сигнализации и готовности инженерных систем по доступной документации",
            "object_types": ["cod"],
            "work_items": ["Контроль доступа", "Пожарная сигнализация", "Актуальность аварийных контактов", "Проверка журналов обслуживания"],
            "periodicity": "semi_annual",
            "responsible_role": "engineer",
            "work_direction": "network_engineer",
        },
        # ========== АПК (камеры) — монтаж / инженер ==========
        {
            "code": "7.5",
            "name": "Юстировка камер",
            "description": "Юстировка (настройка угла обзора) камер видеонаблюдения",
            "object_types": ["camera_integration", "ovn"],
            "work_items": ["Юстировка камер"],
            "periodicity": "monthly",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
        {
            "code": "7.6",
            "name": "Проверка надёжности крепления камер",
            "description": "Проверка крепления камер",
            "object_types": ["camera_integration", "ovn"],
            "work_items": ["Проверка крепления камер"],
            "periodicity": "monthly",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
        {
            "code": "7.7",
            "name": "Проверка крепления коммутационного шкафа",
            "description": "Проверка крепления шкафа",
            "object_types": ["camera_integration", "ovn"],
            "work_items": ["Проверка крепления шкафа"],
            "periodicity": "monthly",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
        {
            "code": "7.8",
            "name": "Обслуживание кабельных линий UTP",
            "description": "Проверка провисов, укладка запасов UTP",
            "object_types": ["camera_integration", "ovn"],
            "work_items": ["Проверка провисов UTP", "Укладка запасов"],
            "periodicity": "monthly",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
        {
            "code": "11.1-11.6",
            "name": "ППР ВОЛС",
            "description": "Планово-профилактические работы на ВОЛС",
            "object_types": ["voles"],
            "work_items": ["11.1", "11.2", "11.3", "11.4", "11.5", "11.6"],
            "periodicity": "twice_year",
            "responsible_role": "engineer",
            "work_direction": "installer",
        },
    ]

    added = 0
    updated = 0
    for t in templates:
        existing = db.query(PPRTemplate).filter(PPRTemplate.code == t["code"]).first()
        if not existing:
            db.add(PPRTemplate(**t))
            added += 1
        else:
            # обновить направление и описание при ресиде
            for k, v in t.items():
                if k != "code":
                    setattr(existing, k, v)
            updated += 1
    db.commit()
    print(f"✓ Шаблоны ППР: +{added}, обновлено {updated}, всего кодов {len(templates)}")



def seed_directories(db: Session):
    """Справочники: подразделения, должности, системные роли (без сотрудников)."""
    roles = [
        {"code": "admin", "name": "Администратор", "description": "Полный доступ к системе"},
        {"code": "head", "name": "Руководитель", "description": "Просмотр, планы, отчёты, дежурства"},
        {"code": "engineer", "name": "Инженер / исполнитель", "description": "Свои задачи ППР, ежедневный контроль"},
        {"code": "viewer", "name": "Наблюдатель", "description": "Только просмотр"},
    ]
    for r in roles:
        if not db.query(SystemRole).filter(SystemRole.code == r["code"]).first():
            db.add(SystemRole(**r))

    # Подразделения и должности из штатной расстановки
    structure = [
        ("Администрация филиала", [
            "Директор филиала", "Главный инженер", "Бухгалтер", "Главный бухгалтер филиала",
            "Специалист по кадрам", "Кладовщик", "Водитель", "Инженер по ОТ и ТБ",
        ], False, False),
        ("Центр мониторинга", [
            "Руководитель центра мониторинга", "Специалист по мониторингу",
        ], False, False),
        ("Отдел доступов и разрешений", [
            "Старший менеджер доступа", "Менеджер доступа",
        ], False, False),
        ("Отдел сетей связи и группа IT", [
            "Начальник отдела сетей связи", "Ведущий сетевой инженер", "Сетевой инженер",
            "Системный администратор",
        ], True, True),
        ("Отдел технической эксплуатации", [
            "Начальник отдела технической эксплуатации", "Инженер отдела технической эксплуатации",
        ], False, False),
        ("Производственный отдел", [
            "Начальник производственного отдела", "Бригадир", "Старший энергетик", "Монтажник",
            "Водитель-оператор автогидроподъемника", "Сварщик ВОЛС",
        ], True, True),
        ("ГПХ", ["Уборщица"], False, False),
    ]

    # Должности, для которых ППР/дежурства по умолчанию
    ppr_positions = {
        "Начальник отдела сетей связи", "Ведущий сетевой инженер", "Сетевой инженер",
        "Системный администратор", "Бригадир", "Монтажник", "Сварщик ВОЛС",
        "Начальник производственного отдела",
    }

    for i, (dept_name, positions, _, _) in enumerate(structure):
        dept = db.query(Department).filter(Department.name == dept_name).first()
        if not dept:
            dept = Department(name=dept_name, sort_order=i * 10, is_active=True)
            db.add(dept)
            db.flush()
        for pname in positions:
            pname = pname.strip()
            exists = db.query(Position).filter(
                Position.name == pname, Position.department_id == dept.id
            ).first()
            if not exists:
                db.add(Position(
                    name=pname,
                    department_id=dept.id,
                    can_do_ppr_default=pname in ppr_positions,
                    can_duty_default=pname in ppr_positions,
                    is_active=True,
                ))
    db.commit()
    print(f"✓ Подразделения: {db.query(Department).count()}, должности: {db.query(Position).count()}, роли: {db.query(SystemRole).count()}")


def seed_users(db: Session):
    """Только учётная запись администратора. Сотрудников добавляют вручную."""
    if not db.query(User).filter(User.username == "admin").first():
        db.add(User(
            username="admin",
            full_name="Администратор",
            role="admin",
            position="Начальник ОСС",
            hashed_password=get_password_hash("admin"),
            can_do_ppr=True,
            can_duty=True,
            is_active=True,
        ))
        db.commit()
        print("✓ Создан admin / admin")
    else:
        print("✓ admin уже есть")



def import_integration_objects(db: Session):
    """Импорт объектов ИНТ из отчёта"""
    xlsx_path = ATTACHMENTS / "Отчет_ППР_АПК_СЕНТЯБРЬ_2026г..xlsx"
    if not xlsx_path.exists():
        print("⚠ Файл отчёта не найден")
        return

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["ППР (ИНТЕГ),"]

    count = 0
    for row in ws.iter_rows(min_row=4, max_row=300, max_col=9, values_only=True):
        num, code, address, cameras, responsible, ppr_date, period, status, work_types = row[:9]
        if not code or not str(code).startswith("ИНТ"):
            continue
        code = str(code).strip()
        if db.query(Object).filter(Object.code == code).first():
            continue

        try:
            cams = int(cameras) if cameras and str(cameras).isdigit() else 0
        except Exception:
            cams = 0

        obj = Object(
            code=code,
            name=str(address)[:100] if address else code,
            address=str(address) if address else None,
            object_type="camera_integration",
            cameras_count=cams,
            status="active",
            responsible_default=str(responsible) if responsible else None,
            notes=None,
        )
        db.add(obj)
        count += 1

        # Если есть дата и статус — создаём запись плана
        if ppr_date and status:
            db.flush()
            plan = PPRPlan(
                object_id=obj.id,
                title=f"ППР {code}",
                planned_date=ppr_date.date() if hasattr(ppr_date, "date") else date(2025, 12, 1),
                fact_date=ppr_date.date() if hasattr(ppr_date, "date") else None,
                responsible=str(responsible) if responsible else None,
                status="done" if "исполнено" in str(status).lower() else "planned",
                work_types_done=str(work_types).split(", ") if work_types else [],
            )
            db.add(plan)

    db.commit()
    print(f"✓ Импортировано объектов ИНТ: {count}")


def import_voles_objects(db: Session):
    """Импорт ВОЛС"""
    xlsx_path = ATTACHMENTS / "Отчет_ППР_АПК_СЕНТЯБРЬ_2026г..xlsx"
    if not xlsx_path.exists():
        return

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "ППР (ВОЛС) г.Ука" not in wb.sheetnames:
        return
    ws = wb["ППР (ВОЛС) г.Ука"]

    count = 0
    for row in ws.iter_rows(min_row=4, max_row=260, max_col=7, values_only=True):
        num, address, responsible, ppr_date, period, status, note = row[:7]
        if not address:
            continue
        code = f"ВОЛС-{num}" if num else f"ВОЛС-{count+1}"
        if db.query(Object).filter(Object.code == code).first():
            continue

        obj = Object(
            code=code,
            name=str(address)[:100],
            address=str(address),
            object_type="voles",
            cameras_count=0,
            status="active",
            responsible_default=str(responsible) if responsible else None,
        )
        db.add(obj)
        count += 1

        if ppr_date and status:
            db.flush()
            plan = PPRPlan(
                object_id=obj.id,
                title=f"ППР ВОЛС {address}",
                planned_date=ppr_date.date() if hasattr(ppr_date, "date") else date(2025, 12, 18),
                fact_date=ppr_date.date() if hasattr(ppr_date, "date") else None,
                responsible=str(responsible) if responsible else None,
                status="done" if "исполнено" in str(status).lower() else "planned",
                work_types_done=["11.1", "11.2", "11.3", "11.4", "11.5", "11.6"],
            )
            db.add(plan)

    db.commit()
    print(f"✓ Импортировано объектов ВОЛС: {count}")


def seed_sample_plans(db: Session):
    """Создаём несколько плановых задач на октябрь 2026 для демонстрации"""
    objects = db.query(Object).filter(Object.object_type == "camera_integration").limit(15).all()
    if not objects:
        return

    templates = {t.code: t for t in db.query(PPRTemplate).all()}
    base_date = date(2026, 10, 15)

    for i, obj in enumerate(objects):
        planned = base_date + timedelta(days=i % 10)
        plan = PPRPlan(
            object_id=obj.id,
            template_id=templates.get("7.5").id if "7.5" in templates else None,
            title=f"ППР {obj.code} — октябрь 2026",
            planned_date=planned,
            responsible=obj.responsible_default or "Инженер ОСС",
            status="planned",
            work_types_done=[],
        )
        db.add(plan)
    db.commit()
    print(f"✓ Создано плановых задач на октябрь: {len(objects)}")


def run_seed():
    create_tables()
    db = SessionLocal()
    try:
        seed_templates(db)
        seed_directories(db)
        seed_users(db)
        import_integration_objects(db)
        import_voles_objects(db)
        seed_sample_plans(db)
        print("\n✅ База данных инициализирована успешно")
    finally:
        db.close()


if __name__ == "__main__":
    run_seed()
