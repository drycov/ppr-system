"""Экспорт отчёта ППР в формате текущего Excel"""
from io import BytesIO
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, NamedStyle
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from models import Object, PPRPlan


THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(bold=True, size=12)
DONE_FILL = PatternFill("solid", fgColor="D1FAE5")
PLANNED_FILL = PatternFill("solid", fgColor="FEF3C7")
OVERDUE_FILL = PatternFill("solid", fgColor="FEE2E2")


def _style_header(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN


def export_ppr_report(
    db: Session,
    year: int,
    month: int,
    object_type: str = None,
) -> bytes:
    """
    Формирует Excel-отчёт в стиле «Отчет_ППР_АПК».
    Листы: ИНТЕГ, ВОЛС (и при необходимости другие).
    """
    wb = Workbook()

    # ----- Лист ИНТЕГ -----
    ws = wb.active
    ws.title = "ППР (ИНТЕГ)"

    month_names = {
        1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель",
        5: "Май", 6: "Июнь", 7: "Июль", 8: "Август",
        9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь",
    }
    title = f"ППР АПК ИНТЕГРАЦИЯ по г.Усть-Каменогорск {month_names.get(month, '')} {year} год"
    ws.merge_cells("A1:I1")
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = [
        "№", "№ ИНТ", "Адрес установки", "Кол-во камер интеграции",
        "Ответственный за ППР", "Дата ППР", "Периодичность",
        "Статус/Примечание", "Виды работы",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=3, column=col, value=h)
    _style_header(ws, 3, len(headers))
    ws.row_dimensions[3].height = 30

    # Объекты ИНТ + их планы за месяц
    month_start = date(year, month, 1)
    if month == 12:
        month_end = date(year + 1, 1, 1)
    else:
        month_end = date(year, month + 1, 1)

    objects = db.query(Object).filter(
        Object.object_type == "camera_integration",
        Object.status == "active",
    ).order_by(Object.code).all()

    row = 4
    for idx, obj in enumerate(objects, 1):
        plan = db.query(PPRPlan).filter(
            PPRPlan.object_id == obj.id,
            PPRPlan.planned_date >= month_start,
            PPRPlan.planned_date < month_end,
        ).order_by(PPRPlan.planned_date.desc()).first()

        # Если нет плана в этом месяце — ищем любой done
        if not plan:
            plan = db.query(PPRPlan).filter(
                PPRPlan.object_id == obj.id,
                PPRPlan.status == "done",
            ).order_by(PPRPlan.fact_date.desc()).first()

        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=obj.code)
        ws.cell(row=row, column=3, value=obj.address or obj.name or "")
        ws.cell(row=row, column=4, value=obj.cameras_count or "")
        ws.cell(row=row, column=5, value=(plan.responsible if plan else obj.responsible_default) or "")
        if plan and (plan.fact_date or plan.planned_date):
            d = plan.fact_date or plan.planned_date
            ws.cell(row=row, column=6, value=d.strftime("%Y-%m-%d") if d else "")
        else:
            ws.cell(row=row, column=6, value="")
        ws.cell(row=row, column=7, value="ежемесячно")
        status = plan.status if plan else ""
        status_ru = {"done": "исполнено", "planned": "запланировано", "overdue": "просрочено", "in_progress": "в работе"}.get(status, status)
        ws.cell(row=row, column=8, value=status_ru)
        works = ", ".join(plan.work_types_done) if plan and plan.work_types_done else ""
        if plan and plan.status == "done" and not works:
            works = "7.5, 7.6, 7.7, 7.8."
        ws.cell(row=row, column=9, value=works)

        fill = None
        if status == "done":
            fill = DONE_FILL
        elif status == "overdue":
            fill = OVERDUE_FILL
        elif status == "planned":
            fill = PLANNED_FILL
        for col in range(1, 10):
            cell = ws.cell(row=row, column=col)
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if fill:
                cell.fill = fill
        row += 1

    # Ширины колонок
    widths = [6, 12, 50, 12, 28, 12, 14, 16, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Примечание
    row += 2
    ws.cell(row=row, column=1, value="Примечание:")
    ws.cell(row=row, column=1).font = Font(bold=True)
    row += 1
    notes = [
        "7.5 — Юстировка камер.",
        "7.6 — Проверка надёжности крепления камер.",
        "7.7 — Проверка надёжности крепления коммутационного шкафа.",
        "7.8 — Обслуживание кабельных линий UTP, проверка провисов, укладка запасов.",
    ]
    for n in notes:
        ws.cell(row=row, column=1, value=n)
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="Составил: Начальник ОСС")
    ws.cell(row=row, column=5, value="Рыков Д.И.")
    row += 2
    ws.cell(row=row, column=1, value="Проверил: Главный инженер")
    ws.cell(row=row, column=5, value="Гениттенов Р.К.")
    row += 2
    ws.cell(row=row, column=1, value="Проверил: Начальник ОТЭ")
    ws.cell(row=row, column=5, value="Ахметвалиев А.Е.")

    # ----- Лист ВОЛС -----
    ws2 = wb.create_sheet("ППР (ВОЛС)")
    ws2.merge_cells("A1:G1")
    ws2["A1"] = f"ППР Волоконно-оптической линии связи (ВОЛС) г.Усть-Каменогорск — {month_names.get(month, '')} {year}"
    ws2["A1"].font = TITLE_FONT

    headers2 = ["№", "Адрес установки", "Ответственный за ППР", "Дата ППР", "Периодичность", "Статус выполнения", "Примечание"]
    for col, h in enumerate(headers2, 1):
        ws2.cell(row=3, column=col, value=h)
    _style_header(ws2, 3, len(headers2))

    voles = db.query(Object).filter(Object.object_type == "voles", Object.status == "active").order_by(Object.code).all()
    row = 4
    for idx, obj in enumerate(voles, 1):
        plan = db.query(PPRPlan).filter(PPRPlan.object_id == obj.id).order_by(PPRPlan.planned_date.desc()).first()
        ws2.cell(row=row, column=1, value=idx)
        ws2.cell(row=row, column=2, value=obj.address or "")
        ws2.cell(row=row, column=3, value=(plan.responsible if plan else obj.responsible_default) or "")
        if plan and (plan.fact_date or plan.planned_date):
            d = plan.fact_date or plan.planned_date
            ws2.cell(row=row, column=4, value=d.strftime("%Y-%m-%d"))
        ws2.cell(row=row, column=5, value="2 раза в год")
        status = plan.status if plan else ""
        status_ru = {"done": "исполнено", "planned": "запланировано", "overdue": "просрочено"}.get(status, status)
        ws2.cell(row=row, column=6, value=status_ru)
        works = ", ".join(plan.work_types_done) if plan and plan.work_types_done else "11.1, 11.2, 11.3, 11.4, 11.5, 11.6."
        ws2.cell(row=row, column=7, value=works if plan and plan.status == "done" else "")
        for col in range(1, 8):
            ws2.cell(row=row, column=col).border = THIN
        row += 1

    for i, w in enumerate([6, 45, 28, 12, 14, 16, 30], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ----- Лист Сводка -----
    ws3 = wb.create_sheet("Сводка")
    ws3["A1"] = f"Сводка выполнения ППР — {month_names.get(month, '')} {year}"
    ws3["A1"].font = TITLE_FONT
    ws3["A3"] = "Показатель"
    ws3["B3"] = "Значение"
    _style_header(ws3, 3, 2)

    total_plans = db.query(PPRPlan).filter(
        PPRPlan.planned_date >= month_start,
        PPRPlan.planned_date < month_end,
    ).count()
    done_plans = db.query(PPRPlan).filter(
        PPRPlan.planned_date >= month_start,
        PPRPlan.planned_date < month_end,
        PPRPlan.status == "done",
    ).count()
    overdue_plans = db.query(PPRPlan).filter(
        PPRPlan.status == "overdue",
    ).count()

    summary = [
        ("Всего объектов ИНТ", len(objects)),
        ("Всего объектов ВОЛС", len(voles)),
        ("Задач ППР за месяц", total_plans),
        ("Исполнено", done_plans),
        ("Просрочено (всего)", overdue_plans),
        ("% выполнения", f"{round(done_plans / total_plans * 100, 1) if total_plans else 0}%"),
    ]
    for i, (k, v) in enumerate(summary, 4):
        ws3.cell(row=i, column=1, value=k).border = THIN
        ws3.cell(row=i, column=2, value=v).border = THIN
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 14

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
