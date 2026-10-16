import sys
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from database import Base
from models import Department, DailyCheck, Object, PPRPlan, PPRTemplate, User
from planning import _planned_dates_for_period, generate_plans_for_period
from report_service import export_ppr_period_report


class PPRCoverageTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine, expire_on_commit=False)
        self.db = self.Session()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def test_daily_periodicity_returns_every_day_of_month(self):
        dates = _planned_dates_for_period(2026, 10, "daily")
        self.assertEqual(len(dates), 31)
        self.assertEqual(dates[0], date(2026, 10, 1))
        self.assertEqual(dates[-1], date(2026, 10, 31))

    def test_daily_tasks_follow_object_responsibility(self):
        department = Department(name="Отдел ОВН", is_active=True, duty_enabled=True)
        owner = User(
            username="ivanov", full_name="Иванов Иван Иванович",
            surname="Иванов", first_name="Иван", patronymic="Иванович",
            department_id=1, is_active=True, can_do_ppr=True,
            ppr_specialty="installer", position="Монтажник",
        )
        other = User(
            username="petrov", full_name="Петров Пётр Петрович",
            department_id=1, is_active=True, can_do_ppr=True,
            ppr_specialty="installer", position="Монтажник",
        )
        obj = Object(
            code="OVN-01", name="Объект ОВН", object_type="ovn",
            status="active", responsible_default="Иванов Иван Иванович",
        )
        template = PPRTemplate(
            code="OVN-D1", name="Ежедневный осмотр ОВН",
            object_types=["ovn"], work_items=["Осмотр"],
            periodicity="daily", work_direction="installer", is_active=True,
            department_ids=[], excluded_department_ids=[],
        )
        self.db.add(department)
        self.db.commit()
        owner.department_id = department.id
        other.department_id = department.id
        self.db.add_all([owner, other, obj, template])
        self.db.commit()

        generate_plans_for_period(
            self.db, 2026, 10, only_without_plan=True,
            prefer_duty=False, include_daily_templates=True,
            department_id=department.id,
        )
        plans = self.db.query(PPRPlan).filter_by(object_id=obj.id).all()
        self.assertEqual(len(plans), 31)
        self.assertTrue(all(plan.assignee_id == owner.id for plan in plans))
        self.assertTrue(all(plan.responsible == owner.full_name for plan in plans))

    def test_daily_template_creates_per_object_per_day_without_duplicates(self):
        department = Department(name="ОСС", is_active=True, duty_enabled=True)
        self.db.add(department)
        obj = Object(code="COD-01", name="ЦОД", object_type="cod", status="active")
        template = PPRTemplate(
            code="DC-1", name="Ежедневный контроль ЦОД",
            object_types=["cod"], work_items=["Температура и климат"],
            periodicity="daily", work_direction="sysadmin", is_active=True,
            department_ids=[], excluded_department_ids=[],
        )
        self.db.add_all([obj, template])
        self.db.commit()

        first = generate_plans_for_period(
            self.db, 2026, 10, only_without_plan=True,
            prefer_duty=False, include_daily_templates=True,
            department_id=department.id,
        )
        self.assertEqual(first["created"], 31)
        self.assertEqual(self.db.query(PPRPlan).filter_by(object_id=obj.id).count(), 31)

        second = generate_plans_for_period(
            self.db, 2026, 10, only_without_plan=True,
            prefer_duty=False, include_daily_templates=True,
            department_id=department.id,
        )
        self.assertEqual(second["created"], 0)
        self.assertEqual(second["skipped"], 31)


    def test_network_department_generates_one_shared_plan_and_report_lists_objects(self):
        department = Department(name="Отдел сетей связи и группа IT", code="NET", is_active=True, duty_enabled=True)
        objects = [
            Object(code="CORE-01", name="Ядро сети", object_type="core_router", status="active"),
            Object(code="SW-01", name="Коммутатор доступа", object_type="access_switch", status="active"),
            Object(code="CAM-01", name="Камера", object_type="camera_integration", status="active"),
        ]
        template = PPRTemplate(
            code="NET-1", name="Проверка сетевой инфраструктуры",
            object_types=["network"], work_items=["Проверить состояние сети"],
            periodicity="monthly", work_direction="network_engineer", is_active=True,
            department_ids=[], excluded_department_ids=[],
        )
        self.db.add_all([department, *objects, template])
        self.db.commit()

        result = generate_plans_for_period(
            self.db, 2026, 10, only_without_plan=True,
            prefer_duty=False, include_daily_templates=True,
            department_id=department.id,
        )
        plans = self.db.query(PPRPlan).filter_by(department_id=department.id).all()
        self.assertEqual(result["created"], 1)
        self.assertEqual(len(plans), 1)
        self.assertIsNone(plans[0].object_id)

        data = export_ppr_period_report(
            self.db, date(2026, 10, 1), date(2026, 10, 31), department_id=department.id
        )
        workbook = load_workbook(BytesIO(data), read_only=True)
        sheet = workbook["Задачи ППР"]
        self.assertEqual(sheet.cell(row=3, column=13).value, "Перечень объектов")
        self.assertEqual(sheet.max_row, 5)
        first_object = sheet.cell(row=4, column=4).value
        second_object = sheet.cell(row=5, column=4).value
        self.assertEqual({first_object, second_object}, {"CORE-01", "SW-01"})
        self.assertEqual(sheet.cell(row=4, column=13).value, "CORE-01 — Ядро сети")
        self.assertEqual(sheet.cell(row=5, column=13).value, "SW-01 — Коммутатор доступа")
        workbook.close()

    def test_report_contains_all_objects_and_daily_checks(self):
        self.db.add_all([
            Object(code="INT-01", name="Камера", object_type="camera_integration", status="active"),
            Object(code="COD-01", name="Дата-центр", object_type="cod", status="active"),
            Object(code="OLD-01", name="Старый объект", object_type="other", status="decommissioned"),
        ])
        self.db.add(DailyCheck(
            check_date=date(2026, 10, 15), engineer="Инженер",
            monitoring_dashboard="ok", critical_services="ok",
            logs_security="ok", config_backups="ok", channel_utilization="ok",
            ups_status="ok", incidents_logged=True, checklist=[],
        ))
        self.db.commit()

        data = export_ppr_period_report(self.db, date(2026, 10, 1), date(2026, 10, 31))
        workbook = load_workbook(BytesIO(data), read_only=True)
        self.assertIn("Все объекты", workbook.sheetnames)
        self.assertIn("Ежедневный контроль", workbook.sheetnames)
        registry = workbook["Все объекты"]
        codes = {registry.cell(row=row, column=2).value for row in range(4, registry.max_row + 1)}
        self.assertEqual(codes, {"INT-01", "COD-01", "OLD-01"})
        self.assertEqual(workbook["Ежедневный контроль"].max_row, 4)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
